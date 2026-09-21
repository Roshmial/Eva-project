#!/usr/bin/env python3
import re
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer

if len(sys.argv) != 2:
    raise SystemExit('usage: audit_xlsx_model.py MODEL.xlsx')
path = Path(sys.argv[1])
wb = load_workbook(path, data_only=False)
valid = set(wb.sheetnames)
missing_sheets = []
external = []
parse_errors = []
blank_single_refs = []
formula_lengths = []

sheet_ref = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-zА-Яа-яЁё0-9_ .&-]+))!\$?([A-Z]{1,3})\$?(\d+)")
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type != 'f' or not isinstance(cell.value, str):
                continue
            formula = cell.value
            formula_lengths.append(len(formula))
            if '[' in formula:
                external.append((ws.title, cell.coordinate, formula))
            try:
                Tokenizer(formula)
            except Exception as exc:
                parse_errors.append((ws.title, cell.coordinate, str(exc)))
            for quoted, plain, col, rownum in sheet_ref.findall(formula):
                sheet = (quoted or plain).replace("''", "'").strip()
                if sheet not in valid:
                    missing_sheets.append((ws.title, cell.coordinate, sheet))
                    continue
                if ':' not in formula and wb[sheet][f'{col}{rownum}'].value is None:
                    blank_single_refs.append((ws.title, cell.coordinate, sheet, f'{col}{rownum}'))

result = {
    'file': str(path),
    'sheets': len(wb.sheetnames),
    'formulas': len(formula_lengths),
    'average_formula_length': round(sum(formula_lengths) / len(formula_lengths), 1) if formula_lengths else 0,
    'maximum_formula_length': max(formula_lengths, default=0),
    'missing_sheet_references': len(missing_sheets),
    'external_links': len(external),
    'formula_parse_errors': len(parse_errors),
    'blank_single_cell_references': len(blank_single_refs),
    'full_calc_on_load': bool(getattr(wb.calculation, 'fullCalcOnLoad', False)),
}
print(result)
if missing_sheets or external or parse_errors or blank_single_refs:
    print({'missing_sheets': missing_sheets[:20], 'external': external[:20], 'parse_errors': parse_errors[:20], 'blank_refs': blank_single_refs[:20]})
    raise SystemExit(1)
