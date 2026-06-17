#!/usr/bin/env python3
"""Best-effort cleanup of unnecessary English business/legal terms in Russian DOCX files.

Usage:
  python scripts/clean_docx_anglicisms_ru.py input.docx [output.docx]

This is intentionally conservative: it keeps product names, company names,
legal/technical abbreviations, and standards, but replaces common consulting
anglicisms in normal text and table cells.
"""
from __future__ import annotations

import argparse
from pathlib import Path
from docx import Document

REPLACEMENTS = {
    "market-entry": "стратегия выхода на рынок",
    "go-to-market": "выход на рынок",
    "enterprise-клиенты": "корпоративные клиенты",
    "enterprise-клиент": "корпоративный клиент",
    "enterprise-контур": "корпоративный контур",
    "enterprise-сегмент": "корпоративный сегмент",
    "enterprise": "корпоративный",
    "security pack": "пакет материалов по безопасности",
    "contract review": "проверка договоров",
    "contract intelligence": "договорная аналитика",
    "contract management": "управление договорами",
    "legal operations": "юридические операции",
    "legal research": "правовое исследование",
    "legal review": "юридическая проверка",
    "legal reasoning": "юридическая аргументация",
    "legal assistant": "юридический помощник",
    "due diligence": "комплексная проверка",
    "drafting": "подготовка документов",
    "redlining": "внесение правок",
    "playbooks": "методики проверки",
    "playbook": "методика проверки",
    "workflow": "рабочий процесс",
    "workflows": "рабочие процессы",
    "deployment": "развёртывание",
    "data residency": "размещение данных",
    "access control": "контроль доступа",
    "audit trail": "журнал аудита",
    "knowledge layer": "слой управления знаниями",
    "document governance": "управление документами",
    "privilege review": "проверка адвокатской тайны и привилегии",
    "private markets": "частные рынки",
    "managed services": "управляемые услуги",
    "in-house legal": "внутренние юридические команды",
    "in-house": "внутренние команды",
    "on-prem": "локальный контур",
    "On-prem": "Локальный контур",
    "No-training": "Запрет обучения на данных клиента",
    "no-training": "запрет обучения на данных клиента",
    "English law evidence": "Подтверждение применимости к английскому праву",
    "Deployment flexibility": "Гибкость развёртывания",
    "Contract risk": "Договорный риск",
    "Research": "Правовое исследование",
    "Doc Q&A": "Вопросы по документам",
    "Summary": "Сводка",
    "Redline": "Правки",
    "Playbook": "Методика проверки",
    "Privilege": "Адвокатская тайна",
    "Residency": "Размещение данных",
    "Access control": "Контроль доступа",
    "Audit": "Аудит",
    "recommended next action": "рекомендуемое следующее действие",
    "jurisdiction warning": "предупреждение по юрисдикции",
    "risks list": "перечень рисков",
    "product discovery": "исследование продукта",
}

RESTORE = {
    "ЮридическийAI": "LegalAI",
    "юридическийAI": "LegalAI",
    "Lexis+ ИИ": "Lexis+ AI",
    "CoCounsel ИИ": "CoCounsel AI",
    "Harvey ИИ": "Harvey AI",
    "Robin ИИ": "Robin AI",
}


def clean_text(text: str) -> str:
    out = text
    for old in sorted(REPLACEMENTS, key=len, reverse=True):
        out = out.replace(old, REPLACEMENTS[old])
    for old, new in RESTORE.items():
        out = out.replace(old, new)
    return out


def set_paragraph_text(paragraph, new_text: str) -> None:
    if paragraph.text == new_text:
        return
    # Keeps paragraph/cell styles but simplifies inline runs.
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
    else:
        paragraph.add_run(new_text)


def clean_docx(input_path: Path, output_path: Path) -> None:
    doc = Document(input_path)
    for p in doc.paragraphs:
        set_paragraph_text(p, clean_text(p.text))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    set_paragraph_text(p, clean_text(p.text))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("output", nargs="?")
    args = ap.parse_args()
    inp = Path(args.input).expanduser().resolve()
    out = Path(args.output).expanduser().resolve() if args.output else inp.with_name(inp.stem + "_ru_polished.docx")
    clean_docx(inp, out)
    print(out)


if __name__ == "__main__":
    main()
