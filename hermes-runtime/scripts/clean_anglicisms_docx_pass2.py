#!/usr/bin/env python3
from pathlib import Path
from docx import Document
import argparse

REPL = {
    'Contract risk': 'Договорный риск',
    'No-training': 'Запрет обучения на данных клиента',
    'English law evidence': 'Подтверждение применимости к английскому праву',
    'English law': 'английское право',
    'English-law': 'по английскому праву',
    'UK/English law': 'право Великобритании / английское право',
    'Deployment flexibility': 'Гибкость развёртывания',
    'customer-managed keys': 'ключи под управлением клиента',
    'private облако': 'частное облако',
    'on-prem': 'локальный контур',
    'On-prem': 'Локальный контур',
    'no-training / retention / deletion': 'запрет обучения / хранение / удаление',
    'no-training': 'запрет обучения на данных клиента',
    'retention': 'хранение данных',
    'deletion': 'удаление данных',
    'legal services': 'юридические услуги',
    'legal документ': 'юридический документ',
    'legal документs': 'юридические документы',
    'legal рабочий процессs': 'юридические рабочие процессы',
    'legal project management': 'управление юридическими проектами',
    'legal operations': 'юридические операции',
    'Legal operations': 'Юридические операции',
    'legal network': 'юридическая сеть',
    'legal productivity': 'продуктивность юридической работы',
    'legal ИИ': 'юридический ИИ',
    'legal intelligence': 'юридическая аналитика',
    'firm intelligence': 'аналитика юридической фирмы',
    'law firm productivity': 'продуктивность юридических фирм',
    'law firm': 'юридическая фирма',
    'multi-jurisdiction research': 'многоюрисдикционное правовое исследование',
    'multi-jurisdiction': 'многоюрисдикционный',
    'document проверка': 'проверка документов',
    'research engine': 'инструмент правового исследования',
    'research': 'правовое исследование',
    'risk layer': 'слой оценки рисков',
    'Data-risk layer': 'слой оценки рисков по данным',
    'данные-risk layer': 'слой оценки рисков по данным',
    'risk': 'риск',
    'knowledge management': 'управление знаниями',
    'knowledge work': 'работа со знаниями',
    'knowledge слой': 'слой знаний',
    'matters': 'дела и проекты',
    'managed content': 'управляемый контент',
    'controlled content': 'управляемый контент',
    'Private markets': 'Частные рынки',
    'private markets': 'частные рынки',
    'high-volume договорs': 'массовые договоры',
    'high-volume': 'массовый',
    'in-house legal': 'внутренние юридические команды',
    'in-house': 'внутренние команды',
    'scaleup/mid-market': 'растущие и средние компании',
    'scaleup': 'растущие компании',
    'mid-market': 'средний рынок',
    'fit': 'соответствие',
    'intelligent договорing': 'интеллектуальная работа с договорами',
    'договорing': 'работа с договорами',
    'procurement/supplier договорing': 'договоры закупок и поставщиков',
    'commercial intelligence': 'коммерческая аналитика',
    'specialist': 'специализированный инструмент',
    'tool': 'инструмент',
    'markets': 'рынки',
    'operations': 'операции',
    'ecosystem': 'экосистема',
    'firm': 'юридическая фирма',
    'product discovery': 'исследование продукта',
    'risks list': 'перечень рисков',
    'jurisdiction warning': 'предупреждение по юрисдикции',
    'recommended next action': 'рекомендуемое следующее действие',
    'governing law': 'применимое право',
    'dispute resolution': 'разрешение споров',
    'product': 'продукт',
    'discovery': 'исследование',
    'jurisdiction': 'юрисдикция',
    'warning': 'предупреждение',
    'recommended': 'рекомендуемый',
    'next action': 'следующее действие',
    'prompts': 'запросов к модели',
    'embeddings': 'векторных представлений',
    'tenant isolation': 'изоляция клиента',
    'model положения договора': 'типовые договорные положения',
    'common law': 'общее право',
    'privilege-preserving процессs': 'процессы с сохранением адвокатской тайны и привилегии',
    'аудит журналы': 'журналы аудита',
    'human проверка': 'проверка человеком',
    'Transfer Limitation Obligation': 'ограничение трансграничной передачи данных',
    'outsourcing': 'аутсорсинг',
    'technology risk': 'технологический риск',
    'MAS-ready пакет': 'пакет соответствия требованиям MAS',
    'DPIA/ИИ impact оценка': 'DPIA / оценка влияния ИИ',
    'immutable аудит журналы': 'неизменяемые журналы аудита',
    'инцидент процесс': 'процесс обработки инцидентов',
    'регион option': 'региональный вариант',
    'deletion контроли': 'контроли удаления данных',
    'regulator/аудит доступ положения договора': 'договорные положения о доступе регулятора и аудита',
    'DPO процесс': 'процесс с участием DPO',
    'ИИ управление документation': 'документация по управлению ИИ',
    'confidentiality safeguards': 'меры защиты конфиденциальности',
    'Data classification контроли': 'контроли классификации данных',
    'change-management': 'управление изменениями',
    'strict конфиденциальность': 'строгая конфиденциальность',
    'договорs': 'договоры',
    'процессs': 'процессы',
    'документs': 'документы',
    'проверка tool': 'инструмент проверки',
}

def clean(s: str) -> str:
    out=s
    # longest first
    for k in sorted(REPL, key=len, reverse=True):
        out=out.replace(k, REPL[k])
    return out

def setp(p, text):
    if p.text == text: return
    for r in p.runs: r.text=''
    if p.runs: p.runs[0].text=text
    else: p.add_run(text)

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('input'); ap.add_argument('output', nargs='?')
    a=ap.parse_args(); inp=Path(a.input); out=Path(a.output) if a.output else inp.with_name(inp.stem+'_clean2.docx')
    d=Document(inp)
    for p in d.paragraphs: setp(p, clean(p.text))
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs: setp(p, clean(p.text))
    d.save(out); print(out)
if __name__=='__main__': main()
